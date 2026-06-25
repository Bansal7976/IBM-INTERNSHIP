import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pptx import Presentation
from pptx.util import Emu

W, H = 13.333, 7.5
EMU = 914400

p = Presentation("Autonomous_Vehicle_Detection_Presentation.pptx")
W_actual = p.slide_width / EMU
H_actual = p.slide_height / EMU
print(f"Slides: {len(p.slides)}, Layout: {W_actual:.2f} x {H_actual:.2f} inches")

for i, slide in enumerate(p.slides, 1):
    print(f"\n=== SLIDE {i} ===")
    # background
    bg = slide.background
    fill = bg.fill
    print(f"  BG fill type: {fill.type}")

    for shp in slide.shapes:
        try:
            l = shp.left / EMU if shp.left else 0
            t = shp.top / EMU if shp.top else 0
            w = shp.width / EMU if shp.width else 0
            h = shp.height / EMU if shp.height else 0
        except: l=t=w=h=0
        name = shp.name
        stype = shp.shape_type
        text = ""
        if shp.has_text_frame:
            for para in shp.text_frame.paragraphs:
                for run in para.runs:
                    text += run.text + " "
            text = text.strip()[:80]
        if shp.has_table:
            print(f"  [TABLE] {name} x={l:.2f} y={t:.2f} w={w:.2f} h={h:.2f}")
            for row in shp.table.rows:
                cells = [c.text.strip() for c in row.cells]
                print(f"    | {' | '.join(cells)}")
        elif text:
            print(f"  [TEXT] {name} x={l:.2f} y={t:.2f} w={w:.2f} h={h:.2f} => '{text}'")
        else:
            print(f"  [SHAPE:{stype}] {name} x={l:.2f} y={t:.2f} w={w:.2f} h={h:.2f}")
